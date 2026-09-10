#!/usr/bin/env python3
"""Create only fictional document inputs for demonstrations and verification."""
from pathlib import Path
import argparse
import csv
from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from pptx import Presentation

ap = argparse.ArgumentParser()
ap.add_argument("out", type=Path)
a = ap.parse_args()
a.out.mkdir(parents=True, exist_ok=True)
doc = Document()
doc.core_properties.author = "DocKit Demo"
doc.core_properties.title = "星河读书会 · 活动说明"
doc.add_heading("星河读书会 · 活动说明", 0)
doc.add_paragraph('本次活动主题是 "读书与日常"。场地面积为 120 平方米,预计有 24 位参与者。')
doc.add_paragraph('联系人说: "请提前 10 分钟到场,带上你喜欢的一本书!"')
doc.add_paragraph('报名网址 https://example.com/signup?day=1 邮箱 hello@example.com 保持原样。')
table = doc.add_table(rows=1, cols=2)
table.rows[0].cells[0].text = "事项"
table.rows[0].cells[1].text = "安排"
for left,right in [("签到", "09:30"), ("分享", '主题: "一本好书"'), ("场地", "120 平方米")]:
    cells=table.add_row().cells;cells[0].text=left;cells[1].text=right
doc.sections[0].header.paragraphs[0].text = '星河读书会 "虚构演示"'
doc.save(a.out / "活动说明.docx")
(a.out / "读书会流程.md").write_text('# 读书会流程\n\n## 签到\n\n09:30 到场。\n\n## 分享\n\n每人分享一本书，约 5 分钟。\n\n## 自由交流\n\n带走一条新想法。\n', encoding="utf-8")
(a.out / "补充说明.md").write_text('# 补充说明\n\n所有文字都是虚构演示，不含真实名单。\n', encoding="utf-8")
(a.out / "报名统计.csv").write_text('活动,人数,场地面积\n读书会,24,120\n观影会,18,90\n', encoding="utf-8")
book=Workbook();book.active.title="读书会";book.active.append(["项目","数量"]);book.active.append(["参加人数",24])
sheet=book.create_sheet("观影会");sheet.append(["项目","数量"]);sheet.append(["参加人数",18])
book.save(a.out / "活动统计.xlsx")
slides=Presentation();slide=slides.slides.add_slide(slides.slide_layouts[1])
slide.shapes.title.text="星河读书会"
slide.placeholders[1].text='活动主题: "读书与日常"\n参与人数: 24'
slides.core_properties.author="DocKit Demo"
slides.save(a.out / "读书会介绍.pptx")
(a.out / "不支持的图片.jpg").write_bytes(b"This is an intentional unsupported-format demonstration, not an image.")
print(a.out)
